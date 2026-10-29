#!/usr/bin/env python3
"""Generate meeting minutes and conservatively update one project issue ledger."""
import argparse
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from model import ContractError, VERSION, canonical, merge, validate, verify_sources
from workbook_io import read_ledger, write_ledger

ROOT=Path(__file__).resolve().parents[1]


def load_profile(path=None):
    profile=json.loads((ROOT/'assets/default-style.json').read_text(encoding='utf-8'))
    base=ROOT/'assets'
    if path is None and (ROOT/'assets/style.local.json').exists():path=ROOT/'assets/style.local.json'
    if path:
        path=Path(path).resolve();profile.update(json.loads(path.read_text(encoding='utf-8')));base=path.parent
    profile['logos']=[str((base/Path(p)).resolve()) for p in profile.get('logos',[])]
    if len(profile['logos'])>2:raise ContractError('At most two header logos are supported')
    for p in profile['logos']:
        if not Path(p).is_file():raise ContractError('Logo file not found: '+p)
    return profile


def set_font(style,profile,size,bold=False):
    style.font.name=profile['font_latin'];style.font.size=Pt(size);style.font.bold=bold
    style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),profile['font_east_asia'])


def word_table(doc,headers,rows,widths):
    table=doc.add_table(rows=1,cols=len(headers));table.autofit=False
    for col,width in zip(table.columns,widths):col.width=Cm(width)
    for c,h in zip(table.rows[0].cells,headers):c.text=h
    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i,row in enumerate([headers,*rows]):
        cells=table.rows[0].cells if i==0 else table.add_row().cells
        for j,(c,value) in enumerate(zip(cells,row)):
            c.width=Cm(widths[j]);c.text=str(value) if value not in (None,'') else '待明确'
            tcpr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'80');el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            if i==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E5EEF4');tcpr.append(fill)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.1
                for run in p.runs:run.bold=i==0
    doc.add_paragraph()
    return table


def write_minutes(data,path,profile):
    d=Document();sec=d.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.left_margin=sec.right_margin=Cm(2);sec.bottom_margin=Cm(1.7)
    sec.header_distance=Cm(.5);sec.footer_distance=Cm(.6)
    set_font(d.styles['Normal'],profile,10.5)
    normal=d.styles['Normal'].paragraph_format;normal.line_spacing=Pt(15);normal.space_after=Pt(5)
    for name,size in [('Title',19),('Heading 1',13),('Heading 2',11),('Heading 3',10.5)]:
        set_font(d.styles[name],profile,size,True)
        d.styles[name].paragraph_format.keep_with_next=True
        d.styles[name].paragraph_format.line_spacing=1.2
    project=data['project'];iv=data.get('interview')
    if not iv:raise ContractError('Minutes require the current interview object')
    header=sec.header.paragraphs[0]
    logos=profile.get('logos',[])
    if logos:
        # Exact body leading clips tall inline logos in header cells.
        header.paragraph_format.line_spacing=Pt(1)
        header.paragraph_format.space_after=Pt(0)
        t=sec.header.add_table(rows=1,cols=2,width=Cm(17))
        t.cell(0,0).paragraphs[0].add_run().add_picture(logos[0],width=Cm(2.25))
        if len(logos)>1:
            p=t.cell(0,1).paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.add_run().add_picture(logos[1],width=Cm(3))
        for cell in t.rows[0].cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing=1.0
                p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(0)
        tail=sec.header.paragraphs[-1]
        tail.paragraph_format.line_spacing=Pt(1)
        tail.paragraph_format.space_after=Pt(0)
    else:header.text=project['name'];header.runs[0].font.size=Pt(8)
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    stage=(' '+iv['stage']) if iv.get('stage') else ''
    run=foot.add_run(project['name']+stage+'  |  ');run.font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');d.settings.element.append(update)
    d.add_paragraph(project['name'],'Title')
    d.add_paragraph(data['minutes'].get('title') or '访谈会议纪要','Title')
    hours=iv.get('duration_hours');duration=f'{hours:g}小时' if isinstance(hours,(int,float)) else None
    day=iv.get('date')
    if day:
        y,m,dd=map(int,day.split('-'));day=f'{y}年{m}月{dd}日'
    word_table(d,['会议事项','内容'],[['会议时间',day],['访谈地点',iv.get('place')],['会议时长',duration],
        ['会议主题',iv.get('topic')],['参会人员','、'.join(iv.get('participants',[])) or None]],[2.8,14.2])
    d.add_paragraph(data['minutes'].get('overview',''))
    numerals=['一','二','三','四','五','六','七','八','九','十']
    def section(part,level,number):
        n=numerals[number-1] if number<=10 else str(number)
        label=f'{n}、' if level==1 else f'（{n}）'
        d.add_heading(label+part['title'],level=min(level,3))
        for j,item in enumerate(part.get('items',[]),1):
            p=d.add_paragraph();p.paragraph_format.left_indent=Cm(.65);p.paragraph_format.first_line_indent=Cm(-.65)
            p.paragraph_format.keep_together=True
            p.add_run(f'{j}.  ');p.add_run(item['label']+'：').bold=True;p.add_run(item['text'])
        for j,sub in enumerate(part.get('subsections',[]),1):section(sub,level+1,j)
    sections=data['minutes'].get('sections',[])
    for i,part in enumerate(sections,1):section(part,1,i)
    actions=data['minutes'].get('actions',[])
    if actions:
        d.add_heading('后续工作安排',1)
        word_table(d,['工作事项','主要内容','牵头方','时间安排'],[
            [a['item']+('' if a.get('agreed') else '（待明确）'),a['content'],a.get('owner'),a.get('when')] for a in actions],[3.1,8.4,2.7,2.8])
    if data['minutes'].get('open_items'):
        d.add_heading('需进一步对齐的事项',1)
        for i,item in enumerate(data['minutes']['open_items'],1):d.add_paragraph(f'{i}. {item}')
    d.core_properties.title=project['name']+' 访谈会议纪要';d.core_properties.author=''
    d.save(path)


def office_convert(source,extension,out_dir):
    exe=shutil.which('soffice') or shutil.which('libreoffice')
    if not exe:raise ContractError('LibreOffice is unavailable; use the environment’s Office tools or omit this option and disclose the unverified step')
    out_dir=Path(out_dir);out_dir.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='interview-office-') as tmp:
        result=subprocess.run([exe,'-env:UserInstallation='+Path(tmp).as_uri(),'--headless','--convert-to',extension,
            '--outdir',str(out_dir),str(Path(source).resolve())],capture_output=True,text=True,timeout=180)
    target=out_dir/(Path(source).stem+'.'+extension)
    if result.returncode or not target.is_file():raise ContractError('Office conversion failed: '+(result.stderr or result.stdout)[-1200:])
    return target


def fingerprint(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest() if Path(path).exists() else None


def review_packet(state):
    pattern=re.compile(r'缺少|缺失|未落实|未开展|不透明|没有|全部|一律|始终|均未')
    by_record={r['id']:[] for r in state['records']}
    for e in state['evidence']:
        for rid in e['record_ids']:
            if rid in by_record:by_record[rid].append({k:e.get(k) for k in ['id','speaker','speaker_role','utterance_type','quote','locator']})
    return {'semantic_review_required':True,'records':[
        {'id':r['id'],'nature':r['nature'],'title':r['title'],'description':r['description'],
         'attention_terms':sorted(set(pattern.findall(r['title']+r['description']))),'evidence':by_record[r['id']]}
        for r in state['records']], 'actions':state.get('minutes',{}).get('actions',[]) if state.get('minutes') else [],
        'instruction':'Review scope, speaker attribution and whether each assertion/owner/date is directly supported. Term flags are prompts, not factual judgments.'}


def build(args):
    data=canonical(json.loads(Path(args.data).read_text(encoding='utf-8')))
    profile=load_profile(args.profile)
    work=Path(args.work_dir).resolve();work.mkdir(parents=True,exist_ok=True)
    targets=[Path(p).resolve() for p in [args.minutes_out,args.ledger_out] if p]
    if len(targets)!=len(set(targets)):raise ContractError('Word and Excel output paths must differ')
    if args.minutes_out and Path(args.minutes_out).exists() and not args.replace_minutes:
        raise ContractError('Minutes already exist; choose a new version or explicitly use --replace-minutes')
    existing_path=Path(args.existing_ledger).resolve() if args.existing_ledger else Path(args.ledger_out).resolve() if args.ledger_out and Path(args.ledger_out).exists() else None
    if args.existing_ledger and not existing_path.is_file():raise ContractError('--existing-ledger does not exist')
    if args.ledger_out and Path(args.ledger_out).exists() and Path(args.ledger_out).resolve()!=existing_path:
        raise ContractError('Existing output differs from --existing-ledger; refusing to overwrite it')
    watched={p:fingerprint(p) for p in set(targets+([existing_path] if existing_path else []))}
    if args.minutes_out:
        if not data.get('minutes') or not data.get('interview'):raise ContractError('Minutes require minutes and interview data')
        data['interview']['minutes_file']=Path(args.minutes_out).name
        for iv in data['interviews']:
            if iv['id']==data['interview']['id']:iv['minutes_file']=Path(args.minutes_out).name
    source_checks=[]
    if data['evidence'] and not args.skip_source_check:
        if not args.sources_dir:raise ContractError('Provide --sources-dir or explicitly use --skip-source-check (draft only)')
        source_checks=verify_sources(data['evidence'],args.sources_dir)
        digests={c['id']:c['sha256'] for c in source_checks}
        for e in data['evidence']:e['source_sha256']=digests[e['id']]
    existing=read_ledger(existing_path) if existing_path else None
    state,changes=merge(existing,data)
    validate(state)
    warnings=[]
    if args.skip_source_check:warnings.append('Source-file quotation checks skipped; draft only')
    report={'version':VERSION,'semantic_review_required':True,'source_checks':source_checks,
        'source_check_scope':'incoming evidence only; historical records preserved', 'changes':changes,
        'record_count':len(state['records']),'warnings':warnings,'outputs':{},'recalculated':False,'previews':[]}
    with tempfile.TemporaryDirectory(dir=work,prefix='build-') as tmp:
        tmp=Path(tmp);staged=[]
        if args.minutes_out:
            target=Path(args.minutes_out).resolve()
            if target.suffix.lower()!='.docx':raise ContractError('Minutes output must be .docx')
            local=tmp/target.name;write_minutes(state,local,profile);staged.append((local,target))
        if args.ledger_out:
            target=Path(args.ledger_out).resolve()
            if target.suffix.lower()!='.xlsx':raise ContractError('Ledger output must be .xlsx (no VBA/macro import)')
            local=tmp/target.name;write_ledger(state,local,profile)
            if args.recalculate:
                converted=office_convert(local,'xlsx',tmp/'recalc');shutil.copy2(converted,local);report['recalculated']=True
            else:warnings.append('Spreadsheet formula caches not recalculated; verify with an Office calculation engine')
            staged.append((local,target))
        from check_deliverables import check_files
        checks=check_files(next((s for s,t in staged if t.suffix=='.docx'),None),
                           next((s for s,t in staged if t.suffix=='.xlsx'),None),require_calculated=args.recalculate)
        report['file_checks']=checks
        if checks['errors']:raise ContractError('\n'.join(checks['errors']))
        expected_logos=len(profile.get('logos',[]))
        if expected_logos:
            if args.minutes_out and checks['details']['minutes']['embedded_images']<expected_logos:raise ContractError('Word logo images were lost')
            if args.ledger_out and any(n<expected_logos for n in checks['details']['ledger']['images_by_sheet'].values()):raise ContractError('Excel logo images were lost')
        if args.preview:
            for local,target in staged:report['previews'].append(str(office_convert(local,'pdf',work/'preview')))
        for p,old_hash in watched.items():
            if fingerprint(p)!=old_hash:raise ContractError('File changed during generation; refusing to overwrite: '+str(p))
        for local,target in staged:
            target.parent.mkdir(parents=True,exist_ok=True)
            if target.exists():
                backup=work/'backups'/datetime.now().strftime('%Y%m%d-%H%M%S-%f')/target.name
                backup.parent.mkdir(parents=True);shutil.copy2(target,backup)
            # Stage beside the destination so replacement is atomic on its filesystem.
            fd,pending=tempfile.mkstemp(prefix='.'+target.stem+'-',suffix=target.suffix,dir=target.parent)
            os.close(fd)
            try:shutil.copy2(local,pending);os.replace(pending,target)
            finally:
                if Path(pending).exists():Path(pending).unlink()
            report['outputs'][target.suffix]=str(target)
    (work/'normalized-data.json').write_text(json.dumps(state,ensure_ascii=False,indent=2),encoding='utf-8')
    (work/'semantic-review.json').write_text(json.dumps(review_packet(state),ensure_ascii=False,indent=2),encoding='utf-8')
    (work/'generation-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    return report


def parser():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--data',required=True);p.add_argument('--minutes-out');p.add_argument('--ledger-out')
    p.add_argument('--existing-ledger');p.add_argument('--profile');p.add_argument('--work-dir',required=True)
    p.add_argument('--sources-dir');p.add_argument('--skip-source-check',action='store_true')
    p.add_argument('--replace-minutes',action='store_true');p.add_argument('--recalculate',action='store_true')
    p.add_argument('--preview',action='store_true')
    return p


if __name__=='__main__':
    p=parser();args=p.parse_args()
    if not (args.minutes_out or args.ledger_out):p.error('At least one output is required')
    try:print(json.dumps(build(args),ensure_ascii=False,indent=2))
    except (ContractError,ValueError,KeyError,OSError) as exc:p.exit(1,str(exc)+'\n')
