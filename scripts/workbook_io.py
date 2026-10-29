"""Canonical Excel storage. Does not silently import legacy/custom workbook schemas."""
from datetime import date, datetime
import math

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.drawing.image import Image
from openpyxl.worksheet.pagebreak import Break

from model import ContractError, VERSION, canonical, split_ids

NAMES = ['问题台账', '核实跟进', '证据明细', '访谈登记', '口径与配置']
HEADERS = [
 ['记录编号','归口小组','协同小组','记录性质','业务领域','问题或事项','具体表现或关注内容','业务影响（含潜在影响）','管理层诉求','涉及职能','核实状态','关联记录编号','来源访谈编号'],
 ['记录编号','问题或事项（自动关联）','核实重点与动作','所需资料或样本','核实完成判据','核实跟进人','计划完成日期','初步原因假设','优先级','核实结论及处置依据','确认人','确认日期','历史编号'],
 ['证据编号','关联记录编号','访谈编号','所属方','发言人或代号','提纲问题编号','来源文件及版本','段落与轮次定位','转写稿原文摘录','发言性质','证据用途','核实提示','来源SHA256'],
 ['访谈编号','阶段','访谈层级','访谈对象','会议日期','访谈地点','会议时长（小时）','参会人员','会议主题','提纲情况','原始转写文件','会议纪要文件']]
WIDTHS = [
 [20,16,26,16,20,36,52,38,38,26,15,25,30],
 [20,36,55,43,57,18,18,38,14,43,18,18,26],
 [22,22,30,15,22,18,45,30,95,18,22,50,30],
 [30,14,20,22,18,25,20,42,44,22,50,55]]
TABLES = ['ProjectIssues','VerificationTasks','EvidenceRecords','InterviewRegistry']
ROLES = {'client':'客户方','consultant':'咨询方','unknown':'未确认'}


def literal(cell, value):
    cell.value = value
    if isinstance(value, str): cell.data_type = 's'  # Source text is never an Excel formula.


def as_date(v):
    return date.fromisoformat(v) if v else None


def iso(v):
    return v.date().isoformat() if isinstance(v, datetime) else v.isoformat() if isinstance(v, date) else v


def read_ledger(path):
    w = openpyxl.load_workbook(path, data_only=False)
    if w.sheetnames != NAMES: raise ContractError('Unsupported workbook: expected v1.2 five-sheet schema; map/backup legacy files first')
    for s, heads in zip(w.worksheets[:4], HEADERS):
        if [s.cell(7,c+1).value for c in range(len(heads))] != heads:
            raise ContractError(f'{s.title}: changed/legacy columns are not imported automatically')
        if any(s.cell(7,c).value is not None for c in range(len(heads)+1,s.max_column+1)):
            raise ContractError(f'{s.title}: extra columns would be lost; explicit migration required')
    config = w['口径与配置']
    if config['F3'].value != VERSION: raise ContractError('Unsupported workbook schema version')
    project = {'name':config['A2'].value, 'code':config['F2'].value, 'groups':[]}
    for row in range(8,config.max_row+1):
        if config.cell(row,1).value:
            project['groups'].append({'name':config.cell(row,1).value,'scope':config.cell(row,2).value or ''})
    def rows(s, width):
        return [list(r[:width]) for r in s.iter_rows(min_row=8,values_only=True) if r[0]]
    follow_rows = rows(w['核实跟进'],13)
    follow = {r[0]:r for r in follow_rows}
    if len(follow) != len(follow_rows): raise ContractError('Duplicate verification task IDs')
    records = []
    for r in rows(w['问题台账'],13):
        if r[0] not in follow: raise ContractError(f'{r[0]}: missing verification task')
        f = follow[r[0]]
        records.append(dict(id=r[0],group=r[1],collaborating_groups=split_ids(r[2]),nature=r[3],domain=r[4],
            title=r[5],description=r[6],impact=r[7],management_request=r[8],departments=split_ids(r[9]),
            status=r[10],related_ids=split_ids(r[11]),interview_ids=split_ids(r[12]),historical_id=f[12],
            verification=dict(focus=f[2],materials=f[3],completion_criteria=f[4],owner=f[5],due_date=iso(f[6]),
                cause_hypothesis=f[7],priority=f[8],conclusion=f[9],confirmer=f[10],confirmed_date=iso(f[11]))))
    if set(follow) != {r['id'] for r in records}: raise ContractError('Orphan verification tasks')
    evidence = []
    for r in rows(w['证据明细'],13):
        evidence.append(dict(id=r[0],record_ids=split_ids(r[1]),interview_id=r[2],
            speaker_role={v:k for k,v in ROLES.items()}.get(r[3],'unknown'),speaker=r[4],outline_question_id=r[5],
            source_file=r[6],locator=r[7],quote=r[8],utterance_type=r[9],purpose=r[10],note=r[11],source_sha256=r[12]))
    interviews = []
    for r in rows(w['访谈登记'],12):
        interviews.append(dict(id=r[0],stage=r[1],level=r[2],object=r[3],date=iso(r[4]),place=r[5],duration_hours=r[6],
            participants=split_ids(r[7]),topic=r[8],has_outline=r[9]=='有提纲',source_file=r[10],minutes_file=r[11]))
    return canonical(dict(project=project,records=records,evidence=evidence,interviews=interviews))


def setup_sheet(s, title, profile, widths):
    s.sheet_view.showGridLines = False
    s.sheet_view.zoomScale = 85
    s.sheet_properties.pageSetUpPr.fitToPage = False
    s.sheet_properties.tabColor = 'C7D8E5'
    for c,width in enumerate(widths,1): s.column_dimensions[openpyxl.utils.get_column_letter(c)].width = width
    s.row_dimensions[1].height = 40; s.row_dimensions[2].height = 29
    literal(s['A2'],title);s['A2'].font = Font(name=profile['font_east_asia'],size=16,bold=True,color='202D3A')
    logos = profile.get('logos', [])
    if logos:
        for path,anchor,width in zip(logos,['A1','D1'],[88,112]):
            img=Image(path);img.height=img.height*width/img.width;img.width=width;s.add_image(img,anchor)
    else:
        literal(s['A1'],'项目工作文件');s['A1'].font=Font(name=profile['font_east_asia'],size=10,color='637686')
    s.freeze_panes='B8'
    s.print_options.horizontalCentered=False
    s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.scale=85
    s.print_title_rows='1:7';s.print_title_cols='A:A'
    s.page_margins.left=s.page_margins.right=.25


def write_table(s, headers, values, widths, name, profile, start_col=1):
    for j,h in enumerate(headers,start_col):
        c=s.cell(7,j);literal(c,h);c.font=Font(name=profile['font_east_asia'],size=10,bold=True,color='243746')
        c.fill=PatternFill('solid',fgColor='E5EEF4');c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    s.row_dimensions[7].height=32
    for row,items in enumerate(values,8):
        lines=1
        for j,value in enumerate(items,start_col):
            c=s.cell(row,j);literal(c,value);c.font=Font(name=profile['font_east_asia'],size=10,color='243746')
            c.alignment=Alignment(vertical='top',wrap_text=True)
            c.fill=PatternFill('solid',fgColor='F4F7F9' if row%2==0 else 'FFFFFF')
            if value is not None:
                n=sum(max(1,math.ceil(sum(2 if ord(ch)>127 else 1 for ch in part)/(widths[j-start_col]*.95))) for part in str(value).split('\n'))
                lines=max(lines,n)
        if lines*15+10>400:raise ContractError(f'{s.title} row {row}: text would be clipped; split the excerpt/item or adjust its layout')
        s.row_dimensions[row].height=max(36,lines*15+10)
    left=openpyxl.utils.get_column_letter(start_col);right=openpyxl.utils.get_column_letter(start_col+len(headers)-1)
    t=Table(displayName=name,ref=f'{left}7:{right}{max(8,7+len(values))}')
    t.tableStyleInfo=TableStyleInfo(name='TableStyleLight1',showRowStripes=False)
    s.add_table(t)
    s.print_area=f'A1:{right}{max(8,7+len(values))}'


def write_ledger(state, path, profile):
    w=openpyxl.Workbook();w.remove(w.active)
    sheets=[w.create_sheet(n) for n in NAMES]
    capacity=max(int(profile.get('capacity',500)),len(state['records'])+100)
    end=7+capacity
    main,follow,ev,registry,config=sheets
    p=state['project']
    titles=[p['name']+' 问题台账','项目问题核实跟进','项目证据明细','项目访谈登记']
    for s,title,widths in zip(sheets[:4],titles,WIDTHS):setup_sheet(s,title,profile,widths)
    mrows=[];frows=[];erows=[];irows=[]
    for r in state['records']:
        v=r['verification']
        mrows.append([r['id'],r['group'],'、'.join(r['collaborating_groups']) or None,r['nature'],r['domain'],r['title'],r['description'],
            r['impact'],r['management_request'],'、'.join(r['departments']) or None,r['status'],'\n'.join(r['related_ids']) or None,'\n'.join(r['interview_ids'])])
        frows.append([r['id'],None,v['focus'],v['materials'],v['completion_criteria'],v['owner'],as_date(v['due_date']),
            v['cause_hypothesis'],v['priority'],v['conclusion'],v['confirmer'],as_date(v['confirmed_date']),r['historical_id']])
    for e in state['evidence']:
        erows.append([e['id'],'\n'.join(e['record_ids']),e['interview_id'],ROLES[e['speaker_role']],e['speaker'],e['outline_question_id'],
            e['source_file'],e['locator'],e['quote'],e['utterance_type'],e['purpose'],e['note'],e['source_sha256']])
    for i in state['interviews']:
        irows.append([i['id'],i['stage'],i['level'],i['object'],as_date(i['date']),i['place'],i['duration_hours'],
            '、'.join(i['participants']),i['topic'],'有提纲' if i['has_outline'] else '无预设提纲',i['source_file'],i['minutes_file']])
    for s,heads,values,widths,table in zip(sheets[:4],HEADERS,[mrows,frows,erows,irows],WIDTHS,TABLES):write_table(s,heads,values,widths,table,profile)
    for s in sheets[:4]:
        s.oddHeader.center.text=p['name'];s.oddHeader.center.size=9
        s.oddFooter.right.text='Page &P / &N'
    for s,columns in [(main,['A','L','M']),(follow,['A','M']),(ev,['A','B','C'])]:
        for c in columns:
            for row in s.iter_rows(min_row=8,max_row=7+len(mrows) if s!=ev else 7+len(erows),min_col=openpyxl.utils.column_index_from_string(c),max_col=openpyxl.utils.column_index_from_string(c)):
                row[0].font=Font(name=profile['font_latin'],size=10,color='243746')
    main.freeze_panes='C8'
    for row in range(8,end+1):
        follow.cell(row,2).value=f'=IF(A{row}="","",INDEX(问题台账!$F$8:$F${end},MATCH(A{row},问题台账!$A$8:$A${end},0)))'
        for c in [7,12]:follow.cell(row,c).number_format='yyyy-mm-dd'
    for row in range(8,8+len(irows)):
        registry.cell(row,5).number_format='yyyy-mm-dd';registry.cell(row,7).number_format='0.0'
    def count(nature,status=None):
        args=f'$A$8:$A${end},"<>",$D$8:$D${end},"{nature}",$K$8:$K${end},"{status or "<>已合并"}"'
        if status is None:args+=f',$K$8:$K${end},"<>已排除"'
        return f'COUNTIFS({args})'
    for a,b,label,formula in [('A','B','问题记录',count('问题')),('C','D','待核实线索',count('待核实线索')),
            ('E','F','其他事项','+'.join(count(n) for n in ['综合议题','管理诉求','实施风险'])),('G','H','已确认问题',count('问题','已确认'))]:
        literal(main[a+'4'],label);main[b+'4']='='+formula;main[b+'4'].font=Font(name=profile['font_latin'],size=12,bold=True)
        main[b+'4'].alignment=Alignment(horizontal='center')
    literal(main['A5'],'项目共用一份台账；按归口小组筛选，小组有效记录数见“口径与配置”。')
    literal(main['A6'],'已合并和已排除记录不计入有效数量；协同小组不重复计数。')
    setup_sheet(config,p['name'],profile,[24,56,30,18,10,18,60,20])
    literal(config['E2'],'项目代号');literal(config['F2'],p['code'])
    literal(config['E3'],'数据版本');literal(config['F3'],VERSION)
    grows=[[g['name'],g.get('scope',''),'按项目配置',None] for g in p['groups']]
    write_table(config,['小组','业务范围','配置说明','有效记录数'],grows,[24,56,30,18],'ProjectGroups',profile)
    for row in range(8,8+len(grows)):
        config.cell(row,4).value=f'=COUNTIFS(问题台账!$A$8:$A${end},"<>",问题台账!$B$8:$B${end},A{row},问题台账!$K$8:$K${end},"<>已合并",问题台账!$K$8:$K${end},"<>已排除")'
    rules=[['登记原则','每场一份纪要，整个项目维护一份台账。','来源按访谈追加'],
        ['归口与协同','每条记录一个归口小组，协同小组可多个。','归口不是业务定责'],
        ['编号','按项目连续编号，不嵌入访谈对象、场次或小组。','归口调整不换编号'],
        ['追加与归并','新访谈补充证据及来源；同一问题不重复创建。','保留已确认信息'],
        ['明确修改','已有业务字段只通过有理由的record_updates修改。','不静默覆盖'],
        ['发言性质','现状、诉求、建议、提问、决定与安排分别记录。','问题不能仅靠提问支持'],
        ['确认与处置','确认、合并和排除须有结论、确认人和日期。','确认不表示整改完成'],
        ['证据','原文逐字摘录，发言性质不等同于真实性验证。','关键疑义继续核实'],
        ['扩展',f'本文件公式、下拉和提示覆盖{capacity}条记录；生成器按数据量扩展。','人工超范围新增需扩展'],
        ['打印','重复项目标题、表头与编号列，长表允许纵向续页。','不要缩成不可读字号']]
    write_table(config,['项目','使用口径','备注'],rules,[18,60,20],'ProjectUsage',profile,start_col=6)
    config.print_area=f'A1:H{max(17,7+len(grows))}';config.freeze_panes=None
    def named(name,ref):w.defined_names.add(DefinedName(name,attr_text=ref))
    for name,ref in [('GroupNames',f"'口径与配置'!$A$8:$A${7+len(grows)}"),('IssueIDs',f"'问题台账'!$A$8:$A${end}"),
        ('CheckIDs',f"'核实跟进'!$A$8:$A${end}"),('CheckConclusions',f"'核实跟进'!$J$8:$J${end}"),
        ('CheckConfirmers',f"'核实跟进'!$K$8:$K${end}"),('CheckDates',f"'核实跟进'!$L$8:$L${end}")]:named(name,ref)
    def dropdown(s,col,formula):
        dv=DataValidation(type='list',formula1=formula,allow_blank=True);dv.errorTitle='请选择有效值';dv.error='请使用下拉选项';dv.showErrorMessage=True;s.add_data_validation(dv);dv.add(f'{col}8:{col}{end}')
    dropdown(main,'B','GroupNames');dropdown(main,'D','"问题,待核实线索,综合议题,管理诉求,实施风险"')
    dropdown(main,'K','"待核实,核实中,已确认,已合并,已排除"');dropdown(follow,'I','"高,中,低,待评估"')
    dropdown(ev,'J','"现状陈述,管理诉求,顾问建议,提问,明确决定,明确安排,其他"')
    red=PatternFill('solid',fgColor='FCE4DF');amber=PatternFill('solid',fgColor='FFF0CA')
    main.conditional_formatting.add(f'A8:A{end}',FormulaRule(formula=[f'AND(A8<>"",COUNTIF($A$8:$A${end},A8)>1)'],fill=red))
    main.conditional_formatting.add(f'B8:B{end}',FormulaRule(formula=['AND($A8<>"",COUNTIF(GroupNames,B8)<>1)'],fill=amber))
    follow.conditional_formatting.add(f'A8:A{end}',FormulaRule(formula=['AND(A8<>"",COUNTIF(IssueIDs,A8)<>1)'],fill=red))
    main.conditional_formatting.add(f'K8:K{end}',FormulaRule(formula=['AND($A8<>"",OR($K8="已确认",$K8="已合并",$K8="已排除"),COUNTIFS(CheckIDs,$A8,CheckConclusions,"<>",CheckConfirmers,"<>",CheckDates,">0")=0)'],fill=amber))
    # Manual horizontal panels keep the title/numbered rows readable when printing.
    for s,split in [(main,7),(follow,5),(ev,8),(registry,7)]:s.col_breaks.append(Break(id=split))
    w.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
    w.save(path)
