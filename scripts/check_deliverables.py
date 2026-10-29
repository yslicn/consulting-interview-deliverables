#!/usr/bin/env python3
"""Check structure/caches; this does not listen to audio or prove semantic truth."""
import argparse
import json
from pathlib import Path
import re
from zipfile import ZipFile
import xml.etree.ElementTree as ET

import openpyxl

from model import ContractError, canonical, counts, validate
from workbook_io import HEADERS, NAMES, read_ledger


def check_files(minutes=None,ledger=None,require_calculated=False,expected=None):
    errors=[];warnings=[];details={}
    if minutes:
        with ZipFile(minutes) as z:
            texts=[]
            for name in z.namelist():
                if name=='word/document.xml' or re.fullmatch(r'word/(header|footer)\d+\.xml',name):
                    root=ET.fromstring(z.read(name));texts.append(''.join(t.text or '' for t in root.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')))
            text='\n'.join(texts)
            if re.search(r'\{\{[^}]+\}\}|TODO|PLACEHOLDER',text,re.I):errors.append('Unfilled Word template placeholders')
            if re.search(r'/Users/|/home/|/mnt/data/|sandbox:/|轮次起点|证据用途|来源SHA256|cite',text):errors.append('Internal provenance/process text leaked into client minutes')
            if re.search(r'\bP\d{3,}\b',text):warnings.append('Word contains P-number-like text: manually check for internal locators or legitimate product codes')
            for label in ['会议时间','访谈地点','会议时长','会议主题','参会人员']:
                if label not in text:errors.append('Missing Word metadata: '+label)
            if '待明确' in text:warnings.append('Word has unresolved metadata/action fields; review before final archiving')
            styles=z.read('word/styles.xml').decode()
            details['minutes']={'visible_characters':len(text),'embedded_images':sum(n.startswith('word/media/') for n in z.namelist()),'font_declarations_present':'eastAsia' in styles}
    if ledger:
        try:
            state=read_ledger(ledger);validate(state)
        except (ContractError,ValueError,KeyError) as exc:
            errors.append(str(exc));return {'errors':errors,'warnings':warnings,'details':details,'semantic_review_required':True}
        formulas=openpyxl.load_workbook(ledger,data_only=False)
        cached=openpyxl.load_workbook(ledger,data_only=True)
        expected_counts=counts(state['records'])
        for col,value in zip(['B','D','F','H'],expected_counts):
            f=formulas['问题台账'][col+'4'];v=cached['问题台账'][col+'4'].value
            if f.data_type!='f':errors.append('Summary formula was overwritten: '+col+'4')
            if v is None:
                (errors if require_calculated else warnings).append('Missing recalculated summary cache: '+col+'4')
            elif v!=value:errors.append(f'Stale summary {col}4: cached {v}, expected {value}')
        titles={r['id']:r['title'] for r in state['records']}
        for row in range(8,8+len(state['records'])):
            rid=formulas['核实跟进'].cell(row,1).value
            if formulas['核实跟进'].cell(row,2).data_type!='f':errors.append(f'{rid}: lookup formula overwritten')
            value=cached['核实跟进'].cell(row,2).value
            if value is None:
                if require_calculated:errors.append(f'{rid}: missing lookup cache')
            elif value!=titles[rid]:errors.append(f'{rid}: lookup title does not match its stable ID')
        for row,g in enumerate(state['project']['groups'],8):
            value=cached['口径与配置'].cell(row,4).value
            expected_group=sum(r['group']==g['name'] and r['status'] not in {'已合并','已排除'} for r in state['records'])
            if value is None:
                (errors if require_calculated else warnings).append('Missing group count cache: '+g['name'])
            elif value!=expected_group:errors.append('Stale group count: '+g['name'])
        for s in formulas:
            if not s.print_area:errors.append(s.title+': no print area')
            if '$A$1' not in str(s.print_area):warnings.append(s.title+': print area may exclude title/branding')
            if s.title!='口径与配置' and (not s.freeze_panes or not s.print_title_cols):errors.append(s.title+': missing freeze/repeated ID column')
            if not s.tables:errors.append(s.title+': native table/filter missing')
            for row in cached[s.title]:
                for c in row:
                    if c.data_type=='e':errors.append(f'{s.title}!{c.coordinate}: cached formula error {c.value}')
        for name in ['问题台账','核实跟进','证据明细']:
            if not formulas[name].data_validations.dataValidation:errors.append(name+': dropdown validation missing')
        if expected:
            expected=canonical(expected)
            for field in ['project','records','evidence','interviews']:
                if state[field]!=expected[field]:errors.append('Workbook differs from normalized working data: '+field)
        details['ledger']={'records':len(state['records']),'interviews':len(state['interviews']),
            'evidence':len(state['evidence']),'groups':len(state['project']['groups']),'counts':expected_counts,
            'images_by_sheet':{s.title:len(s._images) for s in formulas}}
    return {'errors':errors,'warnings':warnings,'details':details,'semantic_review_required':True}


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--minutes');p.add_argument('--ledger');p.add_argument('--expected-data');p.add_argument('--report')
    p.add_argument('--require-calculated',action='store_true')
    args=p.parse_args()
    if not(args.minutes or args.ledger):p.error('Provide at least one file')
    try:
        expected=json.loads(Path(args.expected_data).read_text(encoding='utf-8')) if args.expected_data else None
        result=check_files(args.minutes,args.ledger,args.require_calculated,expected)
    except (OSError,ValueError,KeyError) as exc:result={'errors':[str(exc)],'warnings':[],'semantic_review_required':True}
    result['passed']=not result['errors'];text=json.dumps(result,ensure_ascii=False,indent=2)
    if args.report:Path(args.report).write_text(text,encoding='utf-8')
    print(text);raise SystemExit(0 if result['passed'] else 1)
