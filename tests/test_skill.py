"""All fixtures are fictional. Run with python -m unittest discover -s tests -v."""
from copy import deepcopy
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
import openpyxl
from model import ContractError, canonical, counts, merge, validate, verify_sources
from workbook_io import read_ledger, write_ledger
from build_deliverables import build, load_profile, parser
from check_deliverables import check_files


class SkillTests(unittest.TestCase):
    def setUp(self):
        self.data=canonical(json.loads((ROOT/'examples/interview.json').read_text(encoding='utf-8')))
        self.tmp=tempfile.TemporaryDirectory();self.dir=Path(self.tmp.name)
    def tearDown(self):self.tmp.cleanup()
    def args(self,**overrides):
        f=self.dir/'data.json';f.write_text(json.dumps(self.data,ensure_ascii=False),encoding='utf-8')
        args=parser().parse_args(['--data',str(f),'--sources-dir',str(ROOT/'examples'),
            '--minutes-out',str(self.dir/'minutes.docx'),'--ledger-out',str(self.dir/'master.xlsx'),'--work-dir',str(self.dir/'work')])
        for k,v in overrides.items():setattr(args,k,v)
        return args
    def second_input(self):
        iv=deepcopy(self.data['interviews'][0]);iv['id']='DEMO-INT-002';iv['date']='2026-11-02'
        e=deepcopy(self.data['evidence'][0]);e['id']='DEMO-EVD-005';e['interview_id']=iv['id']
        return canonical({'project':self.data['project'],'interviews':[iv],'evidence':[e],'records':[]})
    def test_example_contract_and_exact_quotes(self):
        validate(self.data);self.assertEqual(counts(self.data['records']),[1,1,0,0])
        self.assertEqual(len(verify_sources(self.data['evidence'],ROOT/'examples')),4)
    def test_wrong_quote_fails(self):
        self.data['evidence'][0]['quote']='This sentence does not exist.'
        with self.assertRaisesRegex(ContractError,'quote not found'):verify_sources(self.data['evidence'],ROOT/'examples')
    def test_source_path_escape_fails(self):
        self.data['evidence'][0]['source_file']='../README.md'
        with self.assertRaisesRegex(ContractError,'escapes'):verify_sources(self.data['evidence'],ROOT/'examples')
    def test_question_is_not_a_confirmed_problem(self):
        self.data['records'][1]['nature']='问题'
        with self.assertRaisesRegex(ContractError,'cannot establish'):validate(self.data)
    def test_consultant_suggestion_is_not_management_request(self):
        self.data['evidence'][2]['speaker_role']='consultant';self.data['evidence'][2]['utterance_type']='顾问建议'
        with self.assertRaisesRegex(ContractError,'management_request'):validate(self.data)
    def test_unagreed_owner_and_date_fail(self):
        self.data['minutes']['actions'][0]['agreed']=False
        with self.assertRaisesRegex(ContractError,'unagreed action'):validate(self.data)
    def test_question_cannot_be_written_as_a_minutes_fact(self):
        self.data['minutes']['sections'][0]['items'][0]['evidence_ids']=['DEMO-EVD-002']
        with self.assertRaisesRegex(ContractError,'fact needs'):validate(self.data)
    def test_large_evidence_is_not_silently_clipped(self):
        self.data['evidence'][0]['quote']='长文本'*1000
        with self.assertRaisesRegex(ContractError,'clipped'):write_ledger(self.data,self.dir/'long.xlsx',load_profile())
    def test_confirmation_needs_real_disposition(self):
        self.data['records'][0]['status']='已确认'
        with self.assertRaisesRegex(ContractError,'confirmation'):validate(self.data)
    def test_duplicate_ids_fail(self):
        self.data['records'].append(deepcopy(self.data['records'][0]))
        with self.assertRaisesRegex(ContractError,'duplicate IDs'):validate(self.data)
    def test_arbitrary_group_count(self):
        for n in [2,5,7]:
            d=deepcopy(self.data)
            d['project']['groups']+= [{'name':f'组{i}','scope':''} for i in range(3,n+1)]
            validate(d);p=self.dir/f'groups-{n}.xlsx';write_ledger(d,p,load_profile())
            w=openpyxl.load_workbook(p)
            self.assertEqual(w.defined_names['GroupNames'].attr_text,f"'口径与配置'!$A$8:$A${7+n}")
    def test_one_project_two_interviews_no_duplicate_issue(self):
        old=deepcopy(self.data);r=old['records'][0];r['status']='已确认'
        r['verification'].update(owner='已指定人员',conclusion='已核对两笔单据',confirmer='测试确认人',confirmed_date='2026-11-02')
        before=deepcopy(old)
        state,log=merge(old,self.second_input());validate(state)
        self.assertEqual(len(state['records']),2);self.assertEqual(len(state['interviews']),2)
        self.assertEqual(state['records'][0]['status'],'已确认')
        self.assertEqual(state['records'][0]['verification']['owner'],'已指定人员')
        self.assertEqual(state['records'][0]['interview_ids'],['DEMO-INT-001','DEMO-INT-002'])
        self.assertEqual(old,before);self.assertFalse(log)
    def test_silent_business_overwrite_fails(self):
        changed=deepcopy(self.data);changed['records'][0]['title']='Changed without an explicit patch'
        with self.assertRaisesRegex(ContractError,'record_updates'):merge(self.data,changed)
    def test_explicit_patch_preserves_id_and_logs_reason(self):
        incoming=self.second_input();incoming['record_updates']=[{'id':'DEMO-ISS-0001','reason':'经项目组复核范围','changes':{'title':'复核后的采购案例范围'}}]
        result,log=merge(self.data,incoming);validate(result)
        self.assertEqual(result['records'][0]['id'],'DEMO-ISS-0001')
        self.assertEqual(log[0]['before']['title'],self.data['records'][0]['title'])
    def test_conflicting_quote_is_not_replaced(self):
        changed=self.second_input();changed['evidence'][0]['id']='DEMO-EVD-001'
        with self.assertRaisesRegex(ContractError,'immutable evidence'):merge(self.data,changed)
    def test_workbook_roundtrip_preserves_business_fields(self):
        p=self.dir/'master.xlsx';write_ledger(self.data,p,load_profile())
        read=read_ledger(p)
        for field in ['project','records','evidence','interviews']:self.assertEqual(read[field],self.data[field],field)
    def test_source_text_cannot_become_a_spreadsheet_formula(self):
        self.data['records'][0]['title']='=1+1'
        p=self.dir/'literal.xlsx';write_ledger(self.data,p,load_profile())
        cell=openpyxl.load_workbook(p)['问题台账']['F8']
        self.assertEqual(cell.data_type,'s');self.assertEqual(cell.value,'=1+1')
    def test_custom_workbook_is_not_silently_flattened(self):
        p=self.dir/'master.xlsx';write_ledger(self.data,p,load_profile());w=openpyxl.load_workbook(p);w.create_sheet('用户补充');w.save(p)
        with self.assertRaisesRegex(ContractError,'Unsupported workbook'):read_ledger(p)
    def test_duplicate_verification_ids_are_not_silently_discarded(self):
        p=self.dir/'duplicate.xlsx';write_ledger(self.data,p,load_profile())
        w=openpyxl.load_workbook(p);w['核实跟进']['A9']=w['核实跟进']['A8'].value;w.save(p)
        with self.assertRaisesRegex(ContractError,'Duplicate verification'):read_ledger(p)
    def test_generator_and_structural_check(self):
        report=build(self.args());self.assertTrue(report['semantic_review_required'])
        result=check_files(self.dir/'minutes.docx',self.dir/'master.xlsx',expected=json.loads((self.dir/'work/normalized-data.json').read_text()))
        self.assertFalse(result['errors']);self.assertTrue(result['warnings']) # no calculation engine was requested
    def test_unrecalculated_file_is_not_certified_as_calculated(self):
        p=self.dir/'master.xlsx';write_ledger(self.data,p,load_profile())
        self.assertTrue(check_files(ledger=p,require_calculated=True)['errors'])
    def test_incremental_build_keeps_master_and_creates_backup(self):
        build(self.args());old=(self.dir/'master.xlsx').read_bytes()
        self.data=self.second_input()
        report=build(self.args(minutes_out=None))
        state=read_ledger(self.dir/'master.xlsx')
        self.assertEqual(len(state['records']),2);self.assertEqual(len(state['interviews']),2)
        backups=list((self.dir/'work/backups').glob('*/master.xlsx'))
        self.assertEqual(len(backups),1);self.assertEqual(backups[0].read_bytes(),old)
    def test_wrong_project_cannot_overwrite_existing_master(self):
        build(self.args());old=(self.dir/'master.xlsx').read_bytes();self.data['project']['code']='OTHER'
        with self.assertRaisesRegex(ContractError,'identity mismatch'):build(self.args(minutes_out=None))
        self.assertEqual((self.dir/'master.xlsx').read_bytes(),old)
    def test_minutes_are_not_overwritten_without_explicit_flag(self):
        build(self.args())
        with self.assertRaisesRegex(ContractError,'Minutes already exist'):build(self.args())


if __name__=='__main__':unittest.main()
